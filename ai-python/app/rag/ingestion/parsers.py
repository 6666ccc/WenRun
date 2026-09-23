"""Document parsers that produce the provider-independent internal model."""

from __future__ import annotations

import importlib.util
import re
from abc import ABC, abstractmethod
from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, BinaryIO, TypeAlias

from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from .models import ElementType, ParsedDocument, ParsedElement

DocumentSource: TypeAlias = bytes | bytearray | str | Path | BinaryIO
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def _source_bytes(source: DocumentSource) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    read = getattr(source, "read", None)
    if not callable(read):
        raise TypeError("source must be bytes, a path, or a readable binary stream")
    value = read()
    return value.encode("utf-8") if isinstance(value, str) else bytes(value)


def _source_name(source: DocumentSource, file_name: str | None) -> str:
    if file_name:
        return Path(file_name).name
    if isinstance(source, (str, Path)):
        return Path(source).name
    raise ValueError("file_name is required when parsing bytes or a stream")


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030")


class DocumentParser(ABC):
    """Boundary between source-format libraries and the ingestion pipeline."""

    name = "abstract"

    @abstractmethod
    def parse(
        self,
        source: DocumentSource,
        *,
        document_id: str,
        file_name: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> ParsedDocument:
        raise NotImplementedError


class NativeDocumentParser(DocumentParser):
    """Offline parser for the formats already supported by the application."""

    name = "native"

    def parse(
        self,
        source: DocumentSource,
        *,
        document_id: str,
        file_name: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> ParsedDocument:
        name = _source_name(source, file_name)
        suffix = Path(name).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(f"unsupported document suffix: {suffix}")
        data = _source_bytes(source)
        if suffix == ".pdf":
            elements = self._pdf_elements(data)
        elif suffix == ".docx":
            elements = self._docx_elements(data)
        elif suffix in {".md", ".markdown"}:
            elements = self._markdown_elements(_decode_text(data))
        else:
            elements = self._text_elements(_decode_text(data))
        return ParsedDocument(
            document_id=document_id,
            file_name=name,
            elements=elements,
            parser=self.name,
            source_type=suffix.lstrip("."),
            metadata=dict(metadata or {}),
        )

    @staticmethod
    def _pdf_elements(data: bytes) -> list[ParsedElement]:
        elements: list[ParsedElement] = []
        for page_number, page in enumerate(PdfReader(BytesIO(data)).pages, start=1):
            for block in re.split(r"\n\s*\n", page.extract_text() or ""):
                if text := block.strip():
                    elements.append(
                        ParsedElement(
                            element_id=f"p{page_number}-e{len(elements) + 1}",
                            element_type=ElementType.PARAGRAPH,
                            text=text,
                            page_number=page_number,
                        )
                    )
        return elements

    @staticmethod
    def _docx_elements(data: bytes) -> list[ParsedElement]:
        document = WordDocument(BytesIO(data))
        elements: list[ParsedElement] = []
        iter_inner_content = getattr(document, "iter_inner_content", None)
        blocks = (
            iter_inner_content()
            if callable(iter_inner_content)
            else [*document.paragraphs, *document.tables]
        )
        for block in blocks:
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                style = (block.style.name if block.style else "").casefold()
                heading_match = re.search(r"heading\s*(\d+)|标题\s*(\d+)", style)
                level = (
                    int(next(value for value in heading_match.groups() if value))
                    if heading_match
                    else None
                )
                if level is not None:
                    element_type = (
                        ElementType.TITLE
                        if level == 1 and not elements
                        else ElementType.HEADING
                    )
                elif "list" in style or "列表" in style:
                    element_type = ElementType.LIST_ITEM
                else:
                    element_type = ElementType.PARAGRAPH
                elements.append(
                    ParsedElement(
                        element_id=f"e{len(elements) + 1}",
                        element_type=element_type,
                        text=text,
                        heading_level=level,
                    )
                )
            elif isinstance(block, Table):
                rows = [
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in block.rows
                ]
                text = "\n".join(row for row in rows if row.strip(" |"))
                if text:
                    elements.append(
                        ParsedElement(
                            element_id=f"e{len(elements) + 1}",
                            element_type=ElementType.TABLE,
                            text=text,
                        )
                    )
        return elements

    @staticmethod
    def _markdown_elements(text: str) -> list[ParsedElement]:
        elements: list[ParsedElement] = []
        blocks = re.split(r"\n\s*\n", text)
        for block in blocks:
            value = block.strip()
            if not value:
                continue
            heading = re.fullmatch(r"(#{1,6})\s+(.+)", value)
            if heading:
                level = len(heading.group(1))
                elements.append(
                    ParsedElement(
                        element_id=f"e{len(elements) + 1}",
                        element_type=ElementType.TITLE if level == 1 and not elements else ElementType.HEADING,
                        text=heading.group(2).strip(),
                        heading_level=level,
                    )
                )
                continue
            lines = value.splitlines()
            if all(re.match(r"^\s*[-*+]\s+", line) for line in lines):
                for line in lines:
                    elements.append(
                        ParsedElement(
                            element_id=f"e{len(elements) + 1}",
                            element_type=ElementType.LIST_ITEM,
                            text=re.sub(r"^\s*[-*+]\s+", "", line).strip(),
                        )
                    )
            else:
                is_table = len(lines) >= 2 and all("|" in line for line in lines[:2])
                elements.append(
                    ParsedElement(
                        element_id=f"e{len(elements) + 1}",
                        element_type=ElementType.TABLE if is_table else ElementType.PARAGRAPH,
                        text=value,
                    )
                )
        return elements

    @staticmethod
    def _text_elements(text: str) -> list[ParsedElement]:
        return [
            ParsedElement(
                element_id=f"e{index}",
                element_type=ElementType.PARAGRAPH,
                text=block.strip(),
            )
            for index, block in enumerate(re.split(r"\n\s*\n", text), start=1)
            if block.strip()
        ]


class DoclingParser(DocumentParser):
    """Optional high-fidelity parser. Importing this module never imports Docling."""

    name = "docling"

    def parse(
        self,
        source: DocumentSource,
        *,
        document_id: str,
        file_name: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> ParsedDocument:
        # Deliberately local: app.main/app.rag.ingest must not load Docling or torch.
        from docling.document_converter import DocumentConverter

        name = _source_name(source, file_name)
        suffix = Path(name).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(f"unsupported document suffix: {suffix}")

        temporary_path: Path | None = None
        source_path: Path
        if isinstance(source, (str, Path)):
            source_path = Path(source)
        else:
            with NamedTemporaryFile(suffix=suffix, delete=False) as temporary:
                temporary.write(_source_bytes(source))
                temporary_path = Path(temporary.name)
                source_path = temporary_path
        try:
            converted = DocumentConverter().convert(source_path)
            elements = list(self._elements(converted.document))
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

        return ParsedDocument(
            document_id=document_id,
            file_name=name,
            elements=elements,
            parser=self.name,
            source_type=suffix.lstrip("."),
            metadata=dict(metadata or {}),
        )

    @staticmethod
    def _elements(document: Any) -> Iterable[ParsedElement]:
        for index, pair in enumerate(document.iterate_items(), start=1):
            item, level = pair if isinstance(pair, tuple) else (pair, None)
            label = str(getattr(item, "label", "")).casefold()
            class_name = type(item).__name__.casefold()
            text = str(getattr(item, "text", "") or "").strip()
            if "table" in class_name and not text:
                export = getattr(item, "export_to_markdown", None)
                if callable(export):
                    text = str(export(document)).strip()
            if not text:
                continue
            if "title" in label:
                element_type = ElementType.TITLE
            elif "section" in label or "heading" in label:
                element_type = ElementType.HEADING
            elif "list" in label:
                element_type = ElementType.LIST_ITEM
            elif "table" in label or "table" in class_name:
                element_type = ElementType.TABLE
            elif "caption" in label:
                element_type = ElementType.CAPTION
            else:
                element_type = ElementType.PARAGRAPH
            provenance = getattr(item, "prov", None) or []
            page_number = getattr(provenance[0], "page_no", None) if provenance else None
            yield ParsedElement(
                element_id=f"e{index}",
                element_type=element_type,
                text=text,
                page_number=page_number,
                heading_level=level if isinstance(level, int) else None,
            )


def default_parser_for(source: DocumentSource, file_name: str | None = None) -> DocumentParser:
    """Prefer Docling for office/PDF files when the optional extra is installed."""

    name = _source_name(source, file_name)
    suffix = Path(name).suffix.lower()
    if suffix in {".pdf", ".docx"} and importlib.util.find_spec("docling") is not None:
        return DoclingParser()
    return NativeDocumentParser()
