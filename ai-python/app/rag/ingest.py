"""资料载入：PDF、Word、TXT、Markdown → Document → 文本片段 → Qdrant。"""

from io import BytesIO
from pathlib import Path
from typing import BinaryIO
from uuid import NAMESPACE_URL, uuid4, uuid5

from app.rag.qdrant import (
    ensure_collection,
    get_embeddings,
    get_qdrant_client,
    get_store,
    hospital_collection,
)
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}


# 步骤一：读取上传文件的二进制内容，并检查扩展名。
def _read_file(file: bytes | bytearray | BinaryIO, filename: str) -> bytes:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"不支持的文件格式：{suffix}，支持：{supported}")

    if isinstance(file, (bytes, bytearray)):
        return bytes(file)

    read = getattr(file, "read", None)
    if not callable(read):
        raise TypeError("file 必须是 bytes 或可读取的二进制文件对象")

    data = read()
    if isinstance(data, str):
        return data.encode("utf-8")
    return bytes(data)


# 步骤二：把不同格式的文件统一转换为 LangChain Document。
# Document 的 page_content 是正文，metadata 保存来源和页码等信息。
def _load_documents(data: bytes, filename: str) -> list[Document]:
    suffix = Path(filename).suffix.lower()
    base_metadata = {
        "source_name": filename,
        "file_type": suffix,
        "document_id": str(uuid4()),
    }

    if suffix == ".pdf":
        reader = PdfReader(BytesIO(data))
        documents = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={**base_metadata, "page": page_number},
                    )
                )
        return documents

    if suffix == ".docx":
        word_document = WordDocument(BytesIO(data))
        paragraphs = [
            paragraph.text.strip()
            for paragraph in word_document.paragraphs
            if paragraph.text and paragraph.text.strip()
        ]
        text = "\n\n".join(paragraphs)
        return [Document(page_content=text, metadata=base_metadata)] if text else []

    # TXT 和 Markdown 都按 UTF-8 优先读取；GB18030 兼容常见中文文本文件。
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")
    text = text.strip()
    return [Document(page_content=text, metadata=base_metadata)] if text else []


# 步骤三：把较长的 Document 切成适合向量检索的较小片段。
# RecursiveCharacterTextSplitter 是 LangChain 对通用文本推荐的切分器。
def _split_documents(documents: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=120,
        add_start_index=True,
    )
    return splitter.split_documents(documents)


# 步骤四：把切好的 Document 写入医院知识库。
def add_hospital_documents(
    documents: list[Document],
    ids: list[str] | None = None,
) -> list[str]:
    if not documents:
        raise ValueError("没有可写入的文档片段")

    client = get_qdrant_client()
    embeddings = get_embeddings()

    # 写入前确保 collection 存在，并且向量维度与 embedding 模型一致。
    if not client.collection_exists(hospital_collection):
        vector_size = len(embeddings.embed_query("dimension probe"))
        ensure_collection(client, hospital_collection, vector_size)

    store = get_store(client, hospital_collection, embeddings)
    return store.add_documents(documents=documents, ids=ids)


# 完整执行一次文件载入：读取、解析、切块、embedding 并写入 Qdrant。
def ingest_file(file: bytes | bytearray | BinaryIO, filename: str) -> dict:
    data = _read_file(file, filename)
    documents = _load_documents(data, filename)
    if not documents:
        raise ValueError("文件没有提取到可用文本；扫描版 PDF 需要先做 OCR")

    chunks = _split_documents(documents)
    if not chunks:
        raise ValueError("文件切分后没有可写入的文本片段")

    # 使用稳定的本次导入 ID，便于后续删除或更新整份文档。
    document_id = str(documents[0].metadata["document_id"])
    # Qdrant 的点 ID 只能使用无符号整数或合法 UUID，不能直接使用 "uuid:序号"。
    chunk_ids = [
        str(uuid5(NAMESPACE_URL, f"{document_id}:{index}"))
        for index in range(len(chunks))
    ]
    written_ids = add_hospital_documents(chunks, ids=chunk_ids)

    return {
        "document_id": document_id,
        "filename": filename,
        "chunk_count": len(written_ids),
    }
