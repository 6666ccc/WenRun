import re
import subprocess
import sys

import pytest
from loguru import logger

from app.rag.ingestion import (
    ChunkingError,
    DocumentCleaner,
    DocumentParseError,
    DocumentType,
    ElementType,
    IngestionConfig,
    IngestionService,
    NativeDocumentParser,
    ParsedDocument,
    ParsedElement,
    RuleBasedDocumentClassifier,
)
from app.rag.ingestion.chunking import ChunkingRouter, HybridChunkingStrategy
from app.rag.ingestion.models import Chunk, QualityStatus
from app.rag.ingestion.parsers import DoclingParser, DocumentParser, default_parser_for
from app.rag.ingestion.pdf_layout import PdfTextLine, elements_from_lines
from app.rag.ingestion.quality import assess_quality


class FakeTokenizer:
    """One Han character is one token; adjacent ASCII letters/digits are one."""

    _parts = re.compile(r"[A-Za-z0-9]+|[\u3400-\u9fff]|[^\s]")

    def encode(self, text, **kwargs):
        return self._parts.findall(text)

    def decode(self, tokens, **kwargs):
        return "".join(tokens)


def _document(*elements: ParsedElement) -> ParsedDocument:
    return ParsedDocument(
        document_id="doc-1",
        file_name="guide.md",
        parser="fake",
        source_type="md",
        elements=list(elements),
    )


def test_markdown_parser_and_cleaner_preserve_structure():
    parsed = NativeDocumentParser().parse(
        "# 门诊须知\n\n## 办理流程\n\n- 挂号\n- 候诊".encode(),
        document_id="doc-1",
        file_name="guide.md",
    )
    cleaned = DocumentCleaner().clean(parsed)

    assert [element.element_type for element in cleaned.elements] == [
        ElementType.TITLE,
        ElementType.HEADING,
        ElementType.LIST_ITEM,
        ElementType.LIST_ITEM,
    ]
    assert cleaned.elements[-1].section_path == ("门诊须知", "办理流程")


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("常见问题\n问：如何预约？\n答：使用小程序。", DocumentType.FAQ),
        ("就诊流程\n第一步：挂号\n第二步：候诊", DocumentType.PROCEDURE),
        ("病历管理制度\n第一条 适用范围", DocumentType.POLICY),
        ("科室目录\n联系电话：123", DocumentType.DIRECTORY),
        ("门诊服务\n挂号流程\n退号规则", DocumentType.HOSPITAL_GUIDE),
        ("医院停车场开放时间说明", DocumentType.GENERAL),
    ],
)
def test_rule_classifier_covers_document_families(text, expected):
    document = _document(
        ParsedElement("e1", ElementType.PARAGRAPH, text),
    )

    assert RuleBasedDocumentClassifier().classify(document) == expected


def test_hybrid_prefers_sentence_boundaries_before_token_limit():
    config = IngestionConfig(max_tokens=10, overlap_tokens=0, min_chunk_tokens=2)
    document = _document(
        ParsedElement(
            "e1",
            ElementType.PARAGRAPH,
            "挂号完成。请按预约时间候诊。检查结束。",
        ),
    )
    strategy = ChunkingRouter(FakeTokenizer(), config).select_strategy(
        DocumentType.GENERAL
    )

    chunks = strategy.split(document)

    assert strategy.name == "hybrid"
    assert [chunk.text for chunk in chunks] == [
        "挂号完成。",
        "请按预约时间候诊。",
        "检查结束。",
    ]
    assert all(len(FakeTokenizer().encode(chunk.text)) <= 10 for chunk in chunks)


def test_single_oversized_sentence_uses_token_window_as_fallback():
    config = IngestionConfig(max_tokens=8, overlap_tokens=2, min_chunk_tokens=2)
    document = _document(
        ParsedElement("e1", ElementType.PARAGRAPH, "一二三四五六七八九十。"),
    )
    strategy = ChunkingRouter(FakeTokenizer(), config).select_strategy(
        DocumentType.GENERAL
    )

    chunks = strategy.split(document)

    assert len(chunks) == 2
    assert all(len(FakeTokenizer().encode(chunk.text)) <= 8 for chunk in chunks)


def test_hospital_guide_keeps_short_sections_in_separate_chunks():
    service = IngestionService(
        parser=NativeDocumentParser(), tokenizer=FakeTokenizer()
    )

    documents = service.ingest(
        (
            "# 门诊服务\n\n"
            "## 挂号流程\n\n请携带身份证在自助机挂号。\n\n"
            "## 退号规则\n\n就诊前可以原路退号。"
        ).encode(),
        document_id="guide-1",
        file_name="门诊指南.md",
    )

    assert len(documents) == 2
    assert {document.metadata["document_type"] for document in documents} == {
        "hospital_guide"
    }
    assert {document.metadata["chunk_strategy"] for document in documents} == {
        "guide_sections"
    }
    assert [document.metadata["section_path"] for document in documents] == [
        "门诊服务 > 挂号流程",
        "门诊服务 > 退号规则",
    ]
    assert "退号规则" not in documents[0].page_content
    assert "挂号流程" not in documents[1].page_content


def test_faq_emits_exactly_one_chunk_per_question_answer_pair():
    service = IngestionService(
        parser=NativeDocumentParser(), tokenizer=FakeTokenizer()
    )

    documents = service.ingest(
        (
            "# 常见问题\n\n"
            "问：如何预约？\n\n答：使用医院小程序。\n\n"
            "问：如何取消？\n\n答：在预约记录中取消。"
        ).encode(),
        document_id="faq-2",
        file_name="门诊FAQ.md",
    )

    assert len(documents) == 2
    assert all(document.metadata["document_type"] == "faq" for document in documents)
    assert all(
        document.metadata["chunk_strategy"] == "faq_pair"
        for document in documents
    )
    assert "如何取消" not in documents[0].page_content
    assert "如何预约" not in documents[1].page_content


def test_medical_paper_headings_select_paper_section_strategy():
    service = IngestionService(
        parser=NativeDocumentParser(), tokenizer=FakeTokenizer()
    )

    documents = service.ingest(
        (
            "# 高血压随访研究\n\n"
            "## 摘要\n\n研究摘要。\n\n"
            "## 讨论\n\n研究讨论。\n\n"
            "## 结论\n\n研究结论。"
        ).encode(),
        document_id="paper-1",
        file_name="随访研究.md",
    )

    assert len(documents) == 3
    assert all(
        document.metadata["document_type"] == "medical_paper"
        for document in documents
    )
    assert [document.metadata["section_path"] for document in documents] == [
        "高血压随访研究 > 摘要",
        "高血压随访研究 > 讨论",
        "高血压随访研究 > 结论",
    ]


def test_service_returns_enriched_langchain_documents_without_external_calls():
    service = IngestionService(
        parser=NativeDocumentParser(),
        tokenizer=FakeTokenizer(),
        config=IngestionConfig(max_tokens=20, overlap_tokens=2),
    )

    documents = service.ingest(
        "# 常见问题\n\n问：如何预约？\n\n答：使用医院小程序。".encode(),
        document_id="faq-1",
        file_name="预约FAQ.md",
        metadata={"version": 3},
    )

    assert documents
    assert documents[0].page_content.startswith("文档：预约FAQ")
    assert documents[0].metadata["document_id"] == "faq-1"
    assert documents[0].metadata["document_type"] == "faq"
    assert documents[0].metadata["chunk_strategy"] == "faq_pair"
    assert documents[0].metadata["version"] == 3


class BrokenParser(DocumentParser):
    name = "broken"

    def parse(self, source, *, document_id, file_name=None, metadata=None):
        detail = "secret body must not be surfaced"
        raise OSError(detail)


def test_service_wraps_parser_failures_in_public_error():
    service = IngestionService(parser=BrokenParser(), tokenizer=FakeTokenizer())
    records = []
    sink_id = logger.add(records.append, format="{message}\n{exception}")
    try:
        with pytest.raises(DocumentParseError) as captured:
            service.ingest(
                b"private content", document_id="bad-1", file_name="bad.txt"
            )
    finally:
        logger.remove(sink_id)

    assert "secret body" not in str(captured.value)
    assert "secret body" not in "".join(records)
    assert "private content" not in "".join(records)


class BrokenTokenizer(FakeTokenizer):
    def encode(self, text, **kwargs):
        raise RuntimeError("tokenizer failed")


def test_service_wraps_chunking_failures():
    service = IngestionService(
        parser=NativeDocumentParser(), tokenizer=BrokenTokenizer()
    )

    with pytest.raises(ChunkingError):
        service.ingest(b"usable text", document_id="bad-2", file_name="bad.txt")


def test_importing_ingestion_does_not_import_docling_or_torch():
    script = (
        "import sys; "
        "import app.main; import app.rag.ingestion; import app.rag.ingest; "
        "import app.rag.chroma; "
        "assert not any(k == 'docling' or k.startswith('docling.') for k in sys.modules); "
        "assert not any(k == 'torch' or k.startswith('torch.') for k in sys.modules)"
    )

    subprocess.run([sys.executable, "-c", script], check=True)


def test_docx_merged_columns_keep_the_grid_width():
    from io import BytesIO

    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=3)
    row = table.rows[0]
    row.cells[0].merge(row.cells[1])
    row.cells[0].text = "合并"
    row.cells[2].text = "单独"
    buffer = BytesIO()
    document.save(buffer)

    elements, _mapping = NativeDocumentParser._parse_docx(buffer.getvalue(), "doc-merge")
    tables = [item for item in elements if item.element_type == ElementType.TABLE]

    assert len(tables) == 1
    assert tables[0].structured_table.rows == [["合并", "合并", "单独"]]


def test_ingest_rejects_a_quality_report_that_cannot_publish():
    service = IngestionService(parser=NativeDocumentParser(), tokenizer=FakeTokenizer())
    content = "| 项目 | 说明 |\n| --- | --- |\n| 挂号 | 携带身份证 |\n".encode()

    with pytest.raises(ValueError, match="质量检查"):
        service.ingest(content, document_id="quality-1", file_name="表.md")


def _docx_bytes(build) -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _mark_header_row(row) -> None:
    from docx.oxml import OxmlElement

    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def _chunks_for(data: bytes, file_name: str, **config: int) -> tuple[list, object]:
    parsed = NativeDocumentParser().parse(data, document_id="sample", file_name=file_name)
    cleaned = DocumentCleaner().clean(parsed)
    settings = {"max_tokens": 512, "overlap_tokens": 64, "min_chunk_tokens": 8}
    settings.update(config)
    strategy = HybridChunkingStrategy(FakeTokenizer(), IngestionConfig(**settings))
    return strategy.split(cleaned), cleaned


def test_word_header_and_footer_stay_out_of_chunks():
    def build(document):
        section = document.sections[0]
        section.header.paragraphs[0].text = "某某医院 内部资料"
        section.footer.paragraphs[0].text = "第 1 页"
        document.add_heading("预约规则", level=1)
        document.add_paragraph("患者可以提前七天预约。")

    chunks, cleaned = _chunks_for(_docx_bytes(build), "预约.docx")
    joined = "\n".join(chunk.text for chunk in chunks)

    assert cleaned.metadata["page_furniture"] == ["某某医院 内部资料", "第 1 页"]
    assert "内部资料" not in joined
    assert "第 1 页" not in joined
    assert "提前七天" in joined


def test_word_table_stays_intact_and_long_tables_repeat_the_header():
    def build_small(document):
        document.add_paragraph("参见下表。")
        table = document.add_table(rows=3, cols=2)
        _mark_header_row(table.rows[0])
        values = [("药品", "剂量"), ("阿司匹林", "100mg"), ("布洛芬", "200mg")]
        for row, cells in zip(table.rows, values, strict=True):
            for cell, value in zip(row.cells, cells, strict=True):
                cell.text = value
        document.add_paragraph("表后仍是正文。")

    chunks, _cleaned = _chunks_for(_docx_bytes(build_small), "剂量.docx")
    table_chunks = [chunk for chunk in chunks if "阿司匹林" in chunk.text]

    assert len(table_chunks) == 1
    assert "布洛芬" in table_chunks[0].text
    assert table_chunks[0].text.startswith("药品 | 剂量")
    assert table_chunks[0].text.count("药品 | 剂量") == 1
    assert all("参见下表" not in chunk.text for chunk in table_chunks)
    assert all("表后仍是正文" not in chunk.text for chunk in table_chunks)

    def build_long(document):
        document.add_paragraph("剂量说明如下。")
        table = document.add_table(rows=6, cols=2)
        _mark_header_row(table.rows[0])
        table.rows[0].cells[0].text = "名"
        table.rows[0].cells[1].text = "量"
        for index, name in enumerate("甲乙丙丁戊", start=1):
            table.rows[index].cells[0].text = name
            table.rows[index].cells[1].text = str(index)

    long_chunks, cleaned = _chunks_for(
        _docx_bytes(build_long),
        "长表.docx",
        max_tokens=6,
        overlap_tokens=0,
        min_chunk_tokens=1,
    )
    parts = [chunk for chunk in long_chunks if chunk.metadata.get("content_kind") == "table"]

    assert len(parts) == 5
    assert all(chunk.text.startswith("名 | 量") for chunk in parts)
    assert all("剂量说明如下" not in chunk.text for chunk in parts)
    report = assess_quality(cleaned, cleaned, len(long_chunks), long_chunks)
    assert any(issue.code == "table_oversized" and issue.severity == "info" for issue in report.issues)
    assert report.status == QualityStatus.PASS


def test_word_formula_stays_with_its_explanation_and_subscripts_are_kept():
    from docx.oxml import parse_xml

    formula = """
    <m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:f>
        <m:num>
          <m:r><m:t>140-age</m:t></m:r>
        </m:num>
        <m:den>
          <m:r><m:t>72</m:t></m:r>
        </m:den>
      </m:f>
    </m:oMath>
    """

    def build(document):
        document.add_paragraph("肌酐清除率计算公式如下：")
        document.add_paragraph()._p.append(parse_xml(formula))
        document.add_paragraph("其中 age 表示年龄。")
        paragraph = document.add_paragraph("溶液为 ")
        paragraph.add_run("H")
        subscript = paragraph.add_run("2")
        subscript.font.subscript = True
        paragraph.add_run("O，钙离子为 ")
        paragraph.add_run("Ca")
        superscript = paragraph.add_run("2+")
        superscript.font.superscript = True
        paragraph.add_run("。")

    parsed = NativeDocumentParser().parse(
        _docx_bytes(build),
        document_id="formula",
        file_name="公式.docx",
    )
    chemical = next(element for element in parsed.elements if "溶液为" in element.text)
    assert "H₂O" in chemical.text
    assert "Ca²⁺" in chemical.text
    assert "H2O" in chemical.search_text
    assert "Ca2+" in chemical.search_text

    cleaned = DocumentCleaner().clean(parsed)
    kept = next(element for element in cleaned.elements if "溶液为" in element.text)
    assert "H₂O" in kept.text
    assert "H2O" in kept.search_text
    formula_element = next(
        element for element in cleaned.elements if element.element_type == ElementType.FORMULA
    )
    assert r"\frac" in formula_element.text
    assert formula_element.metadata["latex_fallback"] is False

    chunks = HybridChunkingStrategy(FakeTokenizer(), IngestionConfig()).split(cleaned)
    formula_chunks = [chunk for chunk in chunks if r"\frac" in chunk.text]
    assert len(formula_chunks) == 1
    assert "肌酐清除率" in formula_chunks[0].text
    assert "年龄" in formula_chunks[0].text


def test_omml_converts_superscript_fraction_and_radical():
    from xml.etree.ElementTree import fromstring

    from app.rag.ingestion.omml import omml_to_latex

    namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    samples = [
        (
            (
                "<m:sSup><m:e><m:r><m:t>height</m:t></m:r></m:e>"
                "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>"
            ),
            "{height}^{2}",
        ),
        (
            (
                "<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>"
                "<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>"
            ),
            r"\frac{a}{b}",
        ),
        (
            (
                '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr>'
                "<m:deg/><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>"
            ),
            r"\sqrt{x}",
        ),
    ]
    for body, expected in samples:
        element = fromstring(f'<m:oMath xmlns:m="{namespace}">{body}</m:oMath>')
        latex, fallback = omml_to_latex(element)
        assert fallback is False
        assert latex == expected


def test_unrecognized_omml_is_marked_for_review():
    from docx.oxml import parse_xml

    formula = """
    <m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:notReal><m:r><m:t>zz</m:t></m:r></m:notReal>
    </m:oMath>
    """

    def build(document):
        document.add_paragraph("无法识别的公式如下：")
        document.add_paragraph()._p.append(parse_xml(formula))

    parsed = NativeDocumentParser().parse(
        _docx_bytes(build),
        document_id="fallback",
        file_name="退回.docx",
    )
    cleaned = DocumentCleaner().clean(parsed)
    chunks = HybridChunkingStrategy(FakeTokenizer(), IngestionConfig()).split(cleaned)
    report = assess_quality(parsed, cleaned, len(chunks), chunks)

    assert any(issue.code == "formula_conversion_fallback" for issue in report.issues)
    assert report.status == QualityStatus.REVIEW


def test_default_parser_stays_native_until_explicitly_configured(monkeypatch):
    monkeypatch.delenv("RAG_PARSER", raising=False)
    assert isinstance(default_parser_for(b"data", file_name="note.pdf"), NativeDocumentParser)

    monkeypatch.setenv("RAG_PARSER", "docling")
    assert isinstance(default_parser_for(b"data", file_name="note.pdf"), DoclingParser)


def test_pdf_layout_removes_repeated_margins_and_merges_pages():
    from io import BytesIO

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    pages = [
        """BT
/F1 9 Tf
72 770 Td
(Wenrun Hospital Internal) Tj
0 -50 Td
/F1 18 Tf
(1.1 Appointment Rules) Tj
0 -40 Td
/F1 12 Tf
(Patients may book seven days ahead.) Tj
0 -560 Td
(The dose continues) Tj
0 -80 Td
/F1 9 Tf
(Page 1) Tj
ET
""",
        """BT
/F1 9 Tf
72 770 Td
(Wenrun Hospital Internal) Tj
0 -80 Td
/F1 12 Tf
(on the following morning.) Tj
0 -660 Td
/F1 9 Tf
(Page 2) Tj
ET
""",
    ]
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for content in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        stream = DecodedStreamObject()
        stream.set_data(content.encode("latin1"))
        page.replace_contents(stream)
    buffer = BytesIO()
    writer.write(buffer)

    parsed = NativeDocumentParser().parse(buffer.getvalue(), document_id="pdf", file_name="rules.pdf")
    cleaned = DocumentCleaner().clean(parsed)
    body = "\n".join(element.text for element in cleaned.elements)

    assert any(element.element_type in {ElementType.TITLE, ElementType.HEADING} for element in cleaned.elements)
    assert any(element.section_path for element in cleaned.elements)
    assert "Wenrun Hospital Internal" not in body
    assert "Page 1" not in body
    assert "Page 2" not in body
    assert any("The dose continues" in element.text and "following morning" in element.text for element in cleaned.elements)
    assert "Wenrun Hospital Internal" in cleaned.metadata["page_furniture"]


def test_numbered_pdf_lines_become_headings_without_a_larger_font():
    lines = [
        PdfTextLine("一、适用范围", 1, 72, 700, 12, 600, 800),
        PdfTextLine("本条说明适用范围，后面还有足够长的正文用来代表普通段落。", 1, 72, 670, 12, 600, 800),
    ]

    elements = elements_from_lines(lines)

    assert elements[0].element_type in {ElementType.TITLE, ElementType.HEADING}
    assert elements[0].text == "一、适用范围"
    assert elements[1].element_type == ElementType.PARAGRAPH


def test_quality_rules_cover_scanned_pdfs_missing_structure_and_residues():
    scanned = ParsedDocument(
        document_id="scan",
        file_name="scan.pdf",
        parser="native",
        source_type="pdf",
        metadata={"pdf_coverage": {"page_count": 4, "text_coverage": 0.25, "missing_text_page_numbers": [2, 3, 4]}},
        elements=[ParsedElement("p", ElementType.PARAGRAPH, "仅第一页有字。", page_number=1)],
    )
    scanned_report = assess_quality(scanned, scanned, 1, [Chunk("仅第一页有字。", ("p",), (1,))])
    assert any(issue.code == "pdf_scanned_document" and issue.blocking for issue in scanned_report.issues)
    assert scanned_report.status == QualityStatus.BLOCKED

    unstructured = ParsedDocument(
        document_id="flat",
        file_name="flat.pdf",
        parser="native",
        source_type="pdf",
        metadata={"pdf_coverage": {"page_count": 6, "text_coverage": 1.0, "missing_text_page_numbers": []}},
        elements=[ParsedElement("p", ElementType.PARAGRAPH, "没有标题的长文档。")],
    )
    unstructured_report = assess_quality(unstructured, unstructured, 1, [Chunk("没有标题的长文档。", ("p",))])
    assert any(issue.code == "pdf_no_structure" for issue in unstructured_report.issues)
    assert unstructured_report.status == QualityStatus.REVIEW

    residue = assess_quality(
        _document(ParsedElement("p", ElementType.PARAGRAPH, "正文")),
        _document(ParsedElement("p", ElementType.PARAGRAPH, "正文")),
        2,
        [
            Chunk("某某医院 内部资料\n第一条。", ("a",)),
            Chunk("某某医院 内部资料\n第二条。", ("b",)),
        ],
    )
    assert any(issue.code == "header_footer_residue" for issue in residue.issues)

    garbled = ParsedDocument(
        document_id="garbled",
        file_name="formula.pdf",
        parser="native",
        source_type="pdf",
        elements=[ParsedElement("p", ElementType.PARAGRAPH, "x = a + b / c * d - e + f")],
    )
    garbled_report = assess_quality(
        garbled,
        garbled,
        1,
        [Chunk("x = a + b / c * d - e + f", ("p",))],
    )
    assert any(issue.code == "formula_garbled_suspected" for issue in garbled_report.issues)

